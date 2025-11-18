"""
Tests for DocumentLoader class.
"""
import os
import pytest
import tempfile
import csv
from pathlib import Path
from unittest.mock import Mock, patch, MagicMock
from io import BytesIO

from src.document_loader import DocumentLoader

# Try importing document creation libraries
try:
    from docx import Document
    from docx.shared import Pt
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from PyPDF2 import PdfWriter, PdfReader
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False


@pytest.fixture
def temp_txt_file():
    """Create a temporary text file."""
    temp_file = tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False)
    temp_file.write("This is a test document.\nWith multiple lines.\n")
    temp_file.close()
    yield temp_file.name
    if os.path.exists(temp_file.name):
        os.unlink(temp_file.name)


@pytest.fixture
def temp_csv_file():
    """Create a temporary CSV file."""
    temp_file = tempfile.NamedTemporaryFile(mode='w', suffix='.csv', delete=False, newline='')
    writer = csv.writer(temp_file)
    writer.writerow(['Name', 'Age', 'City'])
    writer.writerow(['Alice', '30', 'New York'])
    writer.writerow(['Bob', '25', 'San Francisco'])
    writer.writerow(['Charlie', '35', 'Boston'])
    temp_file.close()
    yield temp_file.name
    if os.path.exists(temp_file.name):
        os.unlink(temp_file.name)


@pytest.fixture
def temp_docx_file():
    """Create a temporary DOCX file."""
    if not DOCX_AVAILABLE:
        pytest.skip("python-docx not available")

    temp_file = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
    temp_file.close()

    doc = Document()
    doc.add_paragraph('This is a test DOCX document.')
    doc.add_paragraph('It has multiple paragraphs.')
    doc.add_paragraph('Third paragraph here.')
    doc.save(temp_file.name)

    yield temp_file.name
    if os.path.exists(temp_file.name):
        os.unlink(temp_file.name)


@pytest.fixture
def temp_pdf_file():
    """Create a temporary PDF file using reportlab."""
    if not PDF_AVAILABLE:
        pytest.skip("PDF libraries not available")

    temp_file = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
    temp_file.close()

    # Create PDF with reportlab
    c = canvas.Canvas(temp_file.name, pagesize=letter)
    c.drawString(100, 750, "This is a test PDF document.")
    c.drawString(100, 730, "It has multiple lines.")
    c.drawString(100, 710, "Third line here.")
    c.save()

    yield temp_file.name
    if os.path.exists(temp_file.name):
        os.unlink(temp_file.name)


class TestDocumentLoaderBasics:
    """Basic tests for DocumentLoader class."""

    def test_supported_formats_list(self):
        """Test that supported formats are properly defined."""
        assert hasattr(DocumentLoader, 'SUPPORTED_FORMATS')
        assert isinstance(DocumentLoader.SUPPORTED_FORMATS, list)
        assert '.pdf' in DocumentLoader.SUPPORTED_FORMATS
        assert '.docx' in DocumentLoader.SUPPORTED_FORMATS
        assert '.csv' in DocumentLoader.SUPPORTED_FORMATS

    def test_load_document_nonexistent_file(self):
        """Test loading a nonexistent file raises FileNotFoundError."""
        with pytest.raises(FileNotFoundError):
            DocumentLoader.load_document('/nonexistent/file.pdf')

    def test_load_document_unsupported_format(self, temp_txt_file):
        """Test loading unsupported format raises ValueError."""
        with pytest.raises(ValueError) as excinfo:
            DocumentLoader.load_document(temp_txt_file)

        assert 'Unsupported file format' in str(excinfo.value)

    def test_load_document_returns_dict(self, temp_csv_file):
        """Test that load_document returns a dictionary."""
        result = DocumentLoader.load_document(temp_csv_file)
        assert isinstance(result, dict)
        assert 'content' in result
        assert 'metadata' in result

    def test_metadata_structure(self, temp_csv_file):
        """Test that metadata has correct structure."""
        result = DocumentLoader.load_document(temp_csv_file)
        metadata = result['metadata']

        assert 'source' in metadata
        assert 'filename' in metadata
        assert 'extension' in metadata
        assert 'size' in metadata

        assert metadata['extension'] == '.csv'
        assert metadata['size'] > 0
        assert os.path.basename(temp_csv_file) == metadata['filename']


class TestLoadCSV:
    """Tests for CSV file loading."""

    def test_load_csv_file(self, temp_csv_file):
        """Test loading a CSV file."""
        result = DocumentLoader.load_document(temp_csv_file)

        assert result['content'] != ''
        assert 'Name' in result['content']
        assert 'Alice' in result['content']
        assert 'Bob' in result['content']

    def test_load_csv_metadata(self, temp_csv_file):
        """Test CSV file metadata."""
        result = DocumentLoader.load_document(temp_csv_file)

        assert result['metadata']['extension'] == '.csv'
        assert result['metadata']['filename'].endswith('.csv')

    def test_load_empty_csv(self):
        """Test loading an empty CSV file."""
        temp_file = tempfile.NamedTemporaryFile(mode='w', suffix='.csv', delete=False)
        temp_file.close()

        try:
            result = DocumentLoader.load_document(temp_file.name)
            # Should return something, even if empty
            assert 'content' in result
        finally:
            os.unlink(temp_file.name)

    def test_load_csv_with_special_characters(self):
        """Test loading CSV with special characters."""
        temp_file = tempfile.NamedTemporaryFile(mode='w', suffix='.csv', delete=False, newline='')
        writer = csv.writer(temp_file)
        writer.writerow(['Name', 'Description'])
        writer.writerow(['Test', 'Special chars: @#$%^&*()'])
        writer.writerow(['Unicode', 'Émojis: 😀 🎉'])
        temp_file.close()

        try:
            result = DocumentLoader.load_document(temp_file.name)
            assert 'Special chars' in result['content']
        finally:
            os.unlink(temp_file.name)


class TestLoadDOCX:
    """Tests for DOCX file loading."""

    def test_load_docx_file(self, temp_docx_file):
        """Test loading a DOCX file."""
        result = DocumentLoader.load_document(temp_docx_file)

        assert result['content'] != ''
        assert 'test DOCX document' in result['content']
        assert 'multiple paragraphs' in result['content']

    def test_load_docx_metadata(self, temp_docx_file):
        """Test DOCX file metadata."""
        result = DocumentLoader.load_document(temp_docx_file)

        assert result['metadata']['extension'] == '.docx'
        assert result['metadata']['filename'].endswith('.docx')

    def test_load_docx_preserves_line_breaks(self, temp_docx_file):
        """Test that DOCX loading preserves paragraph breaks."""
        result = DocumentLoader.load_document(temp_docx_file)

        # Should have newlines between paragraphs
        lines = result['content'].split('\n')
        assert len(lines) >= 3


class TestLoadPDF:
    """Tests for PDF file loading."""

    def test_load_pdf_file(self, temp_pdf_file):
        """Test loading a PDF file."""
        result = DocumentLoader.load_document(temp_pdf_file)

        assert result['content'] != ''
        # Check for text content (may vary based on PDF library)
        assert 'test PDF document' in result['content'] or 'PDF' in result['content']

    def test_load_pdf_metadata(self, temp_pdf_file):
        """Test PDF file metadata."""
        result = DocumentLoader.load_document(temp_pdf_file)

        assert result['metadata']['extension'] == '.pdf'
        assert result['metadata']['filename'].endswith('.pdf')
        assert result['metadata']['size'] > 0


class TestLoadDOC:
    """Tests for DOC file loading."""

    def test_load_doc_attempts_docx_loader(self):
        """Test that DOC loader attempts to use DOCX loader."""
        with patch('document_loader.DocumentLoader._load_docx') as mock_docx:
            mock_docx.return_value = "Test content"

            # Create a temp file with .doc extension
            temp_file = tempfile.NamedTemporaryFile(suffix='.doc', delete=False)
            temp_file.write(b"fake doc content")
            temp_file.close()

            try:
                result = DocumentLoader.load_document(temp_file.name)
                assert mock_docx.called
            finally:
                os.unlink(temp_file.name)

    def test_load_doc_invalid_format_raises_error(self):
        """Test that invalid DOC file raises appropriate error."""
        # Create a file with .doc extension but invalid content
        temp_file = tempfile.NamedTemporaryFile(suffix='.doc', delete=False)
        temp_file.write(b"This is not a valid DOC file")
        temp_file.close()

        try:
            with pytest.raises(ValueError) as excinfo:
                DocumentLoader.load_document(temp_file.name)
            assert 'Unable to process .doc file' in str(excinfo.value)
        finally:
            os.unlink(temp_file.name)


class TestLoadODT:
    """Tests for ODT file loading."""

    @patch('document_loader.load_odf')
    @patch('document_loader.teletype.extractText')
    def test_load_odt_calls_correct_methods(self, mock_extract, mock_load):
        """Test that ODT loading calls correct methods."""
        # Create mock document
        mock_doc = MagicMock()
        mock_paragraph = MagicMock()
        mock_doc.getElementsByType.return_value = [mock_paragraph]
        mock_load.return_value = mock_doc
        mock_extract.return_value = "Test paragraph"

        # Create temp ODT file (just for file existence)
        temp_file = tempfile.NamedTemporaryFile(suffix='.odt', delete=False)
        temp_file.close()

        try:
            result = DocumentLoader.load_document(temp_file.name)
            assert mock_load.called
            assert 'Test paragraph' in result['content']
        finally:
            os.unlink(temp_file.name)


class TestDocumentLoaderEdgeCases:
    """Edge case tests for DocumentLoader."""

    def test_load_document_with_uppercase_extension(self):
        """Test loading file with uppercase extension."""
        temp_file = tempfile.NamedTemporaryFile(mode='w', suffix='.CSV', delete=False, newline='')
        writer = csv.writer(temp_file)
        writer.writerow(['Test', 'Data'])
        temp_file.close()

        try:
            result = DocumentLoader.load_document(temp_file.name)
            assert result['metadata']['extension'] == '.csv'  # Should be normalized to lowercase
        finally:
            os.unlink(temp_file.name)

    def test_metadata_absolute_path(self, temp_csv_file):
        """Test that metadata contains absolute path."""
        result = DocumentLoader.load_document(temp_csv_file)

        assert os.path.isabs(result['metadata']['source'])

    def test_load_document_path_with_spaces(self):
        """Test loading document with spaces in filename."""
        temp_dir = tempfile.mkdtemp()
        file_path = os.path.join(temp_dir, 'test file with spaces.csv')

        with open(file_path, 'w', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(['Column1', 'Column2'])
            writer.writerow(['Value1', 'Value2'])

        try:
            result = DocumentLoader.load_document(file_path)
            assert result['content'] != ''
            assert 'test file with spaces.csv' in result['metadata']['filename']
        finally:
            os.unlink(file_path)
            os.rmdir(temp_dir)


class TestStaticMethods:
    """Tests for individual static loader methods."""

    def test_load_csv_static_method(self, temp_csv_file):
        """Test _load_csv static method directly."""
        content = DocumentLoader._load_csv(temp_csv_file)
        assert isinstance(content, str)
        assert 'Alice' in content

    def test_load_docx_static_method(self, temp_docx_file):
        """Test _load_docx static method directly."""
        content = DocumentLoader._load_docx(temp_docx_file)
        assert isinstance(content, str)
        assert len(content) > 0

    def test_load_pdf_static_method(self, temp_pdf_file):
        """Test _load_pdf static method directly."""
        content = DocumentLoader._load_pdf(temp_pdf_file)
        assert isinstance(content, str)
